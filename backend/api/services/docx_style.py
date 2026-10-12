"""
The house style for the Word documents that become PDF reports: fonts that
carry Bangla, shaded header rows, fixed column widths LibreOffice respects.

Shared by scripts/make_production_report_template.py (which builds a
template once) and api.services.expense_report (which builds its document
per request).
"""
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
HEAD_FILL = "EEF2EE"
TOTAL_FILL = "E3ECE5"


def fonts(rpr_parent) -> None:
    """Noto Sans for Latin text, Noto Sans Bengali for Bangla (Word calls it
    a complex script, so it takes the cs font), marked as Bangla so the
    converter shapes it."""
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr, name in (
        ("w:ascii", "Noto Sans"),
        ("w:hAnsi", "Noto Sans"),
        ("w:eastAsia", "Noto Sans"),
        ("w:cs", "Noto Sans Bengali"),
    ):
        rfonts.set(qn(attr), name)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-GB")
    lang.set(qn("w:bidi"), "bn-BD")
    rpr.append(lang)


def cs_size(element, size: float) -> None:
    """The size for complex-script (Bangla) text, which Word keeps apart."""
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(round(size * 2)))
    element.get_or_add_rPr().append(szcs)


def page(doc) -> None:
    """A4 with the report margins, and the Normal and table styles set to
    the house fonts."""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.5)
    cs_size(normal.element, 9.5)
    normal.font.color.rgb = INK
    fonts(normal.element)
    normal.paragraph_format.space_after = Pt(2)
    fonts(doc.styles["Table Grid"].element)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def write(cell, text: str, *, bold=False, size=None, color=None, right=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
        cs_size(run._element, size)
    if color:
        run.font.color.rgb = color


def heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def fixed(t, widths: list[float]) -> None:
    """Column widths LibreOffice keeps: a fixed layout and the grid set,
    not only each cell's preferred width."""
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)


def table(doc, heads: list[str], widths: list[float], right: set[int]):
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed(t, widths)
    for i, (cell, text) in enumerate(zip(t.rows[0].cells, heads)):
        write(cell, text, bold=True, size=8.5, color=MUTED, right=i in right)
        shade(cell, HEAD_FILL)
    for r in t.rows:
        for cell, w in zip(r.cells, widths):
            cell.width = Cm(w)
    return t


def row(t, values: list[str], widths: list[float], right: set[int], *, fill=None, bold=False):
    cells = t.add_row().cells
    for i, (cell, text, w) in enumerate(zip(cells, values, widths)):
        cell.width = Cm(w)
        write(cell, text, bold=bold, right=i in right)
        if fill:
            shade(cell, fill)
    return cells


def keep_rows_whole(t) -> None:
    """Never break a row across two pages: a row moves whole to the next."""
    for r in t.rows:
        el = OxmlElement("w:cantSplit")
        el.set(qn("w:val"), "true")
        r._tr.get_or_add_trPr().append(el)


def repeat_header(t) -> None:
    """Repeat the table's first row at the top of every page it runs onto."""
    trpr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)
