"""
Build the Word template for a production report:
api/reports/production_day.docx.

The template is plain Word with docxtpl tags in it ({{ … }} for a value,
{%tr for … %} rows for a table that grows), so it can be opened and restyled
in Word or LibreOffice later — keep the tags as they are. This script only
makes the first version; it is not run at request time.

Run from backend/:  python -m scripts.make_production_report_template
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "api" / "reports" / "production_day.docx"

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
HEAD_FILL = "EEF2EE"
TOTAL_FILL = "E3ECE5"


def fonts(rpr_parent) -> None:
    """Noto Sans for Latin text, Noto Sans Bengali for Bangla (Word calls it
    a complex script, so it takes the cs font), marked as Bangla so the
    converter shapes it."""
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr, name in (
        ("w:ascii", "Noto Sans"),
        ("w:hAnsi", "Noto Sans"),
        ("w:eastAsia", "Noto Sans"),
        ("w:cs", "Noto Sans Bengali"),
    ):
        rfonts.set(qn(attr), name)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-GB")
    lang.set(qn("w:bidi"), "bn-BD")
    rpr.append(lang)


def cs_size(element, size: float) -> None:
    """The size for complex-script (Bangla) text, which Word keeps apart."""
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(round(size * 2)))
    element.get_or_add_rPr().append(szcs)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def write(cell, text: str, *, bold=False, size=None, color=None, right=False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
        cs_size(run._element, size)
    if color:
        run.font.color.rgb = color


def heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def fixed(t, widths: list[float]) -> None:
    """Column widths LibreOffice keeps: a fixed layout and the grid set,
    not only each cell's preferred width."""
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)


def table(doc, heads: list[str], widths: list[float], right: set[int]):
    t = doc.add_table(rows=1, cols=len(heads))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed(t, widths)
    for i, (cell, text) in enumerate(zip(t.rows[0].cells, heads)):
        write(cell, text, bold=True, size=8.5, color=MUTED, right=i in right)
        shade(cell, HEAD_FILL)
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)
    return t


def row(t, values: list[str], widths: list[float], right: set[int], *, fill=None, bold=False):
    cells = t.add_row().cells
    for i, (cell, text, w) in enumerate(zip(cells, values, widths)):
        cell.width = Cm(w)
        write(cell, text, bold=bold, right=i in right)
        if fill:
            shade(cell, fill)
    return cells


def loop(t, header_tag: str, values: list[str], widths, right, *, end="{%tr endfor %}"):
    """A row per item: the tag rows around the data row vanish when filled."""
    row(t, [header_tag] + [""] * (len(values) - 1), widths, right)
    row(t, values, widths, right)
    row(t, [end] + [""] * (len(values) - 1), widths, right)


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.5)
    cs_size(normal.element, 9.5)
    normal.font.color.rgb = INK
    fonts(normal.element)
    normal.paragraph_format.space_after = Pt(2)
    for name in ("Table Grid",):
        fonts(doc.styles[name].element)

    # --- Title -----------------------------------------------------------------
    top = doc.add_table(rows=1, cols=2)
    fixed(top, [11.8, 6])
    left, rightc = top.rows[0].cells
    left.width, rightc.width = Cm(11.8), Cm(6)
    write(left, "{{ store }}", bold=True, size=16)
    p = left.add_paragraph()
    r = p.add_run("Production report")
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED
    write(rightc, "Date: {{ date }}", bold=True, size=10.5, right=True)
    p = rightc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Production time: {{ time }}")
    r.font.color.rgb = MUTED

    # --- The four figures --------------------------------------------------------
    doc.add_paragraph()
    kpi = doc.add_table(rows=2, cols=4)
    kpi.style = "Table Grid"
    fixed(kpi, [4.45, 4.45, 4.45, 4.45])
    labels = ["Jars produced", "Cooking batches (patils)", "Total production cost", "Cost per jar (avg.)"]
    values = ["{{ jars }}", "{{ patils }}", "{{ total_cost }}", "{{ cost_per_jar }}"]
    for i in range(4):
        write(kpi.rows[0].cells[i], labels[i], size=8.5, color=MUTED)
        write(kpi.rows[1].cells[i], values[i], bold=True, size=14)
        shade(kpi.rows[0].cells[i], HEAD_FILL)
        shade(kpi.rows[1].cells[i], HEAD_FILL)

    # --- Workforce and pay -------------------------------------------------------
    heading(doc, "Workforce & cook salary")
    w, rt = [7.8, 3, 3.5, 3.5], {1, 2, 3}
    t = table(doc, ["Cook", "Count", "Pay each", "Total"], w, rt)
    row(t, ["Male cook", "{{ male_cooks }}", "{{ male_rate }}", "{{ male_total }}"], w, rt)
    row(t, ["Female cook", "{{ female_cooks }}", "{{ female_rate }}", "{{ female_total }}"], w, rt)
    row(t, ["Total labour cost", "{{ total_cooks }}", "", "{{ labour_cost }}"], w, rt, fill=TOTAL_FILL, bold=True)

    heading(doc, "Shift information")
    w, rt = [3, 10.8, 4], {2}
    t = table(doc, ["Shift", "Time", "Cooks"], w, rt)
    loop(t, "{%tr for s in shifts %}", ["{{ s.name }}", "{{ s.time }}", "{{ s.cooks }}"], w, rt)

    # --- Bazar list --------------------------------------------------------------
    heading(doc, "Raw materials (bazar list)")
    w, rt = [1, 7.3, 3.5, 3, 3], {2, 3, 4}
    t = table(doc, ["#", "Item", "Quantity", "Unit price", "Total"], w, rt)
    loop(
        t,
        "{%tr for m in materials %}",
        ["{{ m.no }}", "{{ m.item }}", "{{ m.quantity }}", "{{ m.unit_price }}", "{{ m.amount }}"],
        w,
        rt,
    )
    row(t, ["", "Total raw material cost", "", "", "{{ materials_cost }}"], w, rt, fill=TOTAL_FILL, bold=True)

    # --- What was cooked ---------------------------------------------------------
    heading(doc, "Cooking details")
    w, rt = [8.8, 3, 3, 3], {1, 2, 3}
    t = table(doc, ["Product", "No. of patil", "Jars per patil", "Total jars"], w, rt)
    loop(
        t,
        "{%tr for b in batches %}",
        ["{{ b.product }}", "{{ b.patils }}", "{{ b.per_patil }}", "{{ b.jars }}"],
        w,
        rt,
    )
    row(t, ["Total jars produced", "{{ patils }}", "", "{{ jars }}"], w, rt, fill=TOTAL_FILL, bold=True)

    # --- Other costs ---------------------------------------------------------------
    heading(doc, "Other costs")
    w, rt = [7.8, 6.5, 3.5], {2}
    t = table(doc, ["Cost", "Note", "Amount"], w, rt)
    loop(t, "{%tr for o in other_costs %}", ["{{ o.label }}", "{{ o.note }}", "{{ o.amount }}"], w, rt)

    # --- Where the money went -------------------------------------------------------
    heading(doc, "Total cost breakdown")
    w, rt = [10.8, 4, 3], {1, 2}
    t = table(doc, ["Part", "Amount", "Share"], w, rt)
    loop(t, "{%tr for c in breakdown %}", ["{{ c.label }}", "{{ c.amount }}", "{{ c.share }}"], w, rt)
    row(t, ["Total production cost", "{{ total_cost }}", "100%"], w, rt, fill=TOTAL_FILL, bold=True)
    row(t, ["Cost per jar", "{{ cost_per_jar }}", ""], w, rt, bold=True)

    heading(doc, "Note")
    doc.add_paragraph("{{ note }}")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("Generated {{ generated_at }} by {{ generated_by }}")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
