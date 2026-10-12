"""
Expenses ticked on the Expenses page, as a PDF.

A plain list with a total, so the document is built here with python-docx
in the house style (api.services.docx_style) rather than from a Word
template, then turned into a PDF by LibreOffice the same way production
reports are (api.services.production_report: its throwaway profile, its
limit of two at once).
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from fastapi.concurrency import run_in_threadpool

from api.models import Expense
from api.services import production_report
from api.services.docx_style import (
    MUTED,
    TOTAL_FILL,
    fixed,
    keep_rows_whole,
    page,
    repeat_header,
    row,
    table,
    write,
)
from api.services.production_report import taka

DHAKA = timezone(timedelta(hours=6), "Asia/Dhaka")

# #, date, category, item details, added by, amount: 17.8 cm, the page's width
# inside its margins.
WIDTHS = [0.8, 2.4, 3.0, 6.2, 2.8, 2.6]
RIGHT = {5}


def _date(d) -> str:
    return d.strftime("%d %b %Y").lstrip("0")


def report_context(
    expenses: list[Expense], *, store: str, generated_by: str, now: datetime
) -> dict:
    """Everything the document shows, as text. Pure: expenses in (in the
    order they are to be listed), strings out."""
    days = sorted({e.spent_at.astimezone(DHAKA).date() for e in expenses})
    first, last = days[0], days[-1]
    period = _date(first) if first == last else f"{_date(first)} – {_date(last)}"
    span = first.isoformat() if first == last else f"{first.isoformat()}_{last.isoformat()}"
    return {
        "store": store,
        "period": period,
        "count": len(expenses),
        "rows": [
            {
                "no": str(i + 1),
                "date": _date(e.spent_at.astimezone(DHAKA).date()),
                "category": e.category,
                "item": e.item,
                "note": e.note or "",
                "added_by": e.added_by_name,
                "amount": taka(e.amount),
            }
            for i, e in enumerate(expenses)
        ],
        "total": taka(sum(e.amount for e in expenses)),
        "generated_at": now.strftime("%d %b %Y, %I:%M %p").lstrip("0").replace(" 0", " "),
        "generated_by": generated_by,
        "filename": f"expenses-{span}.pdf",
    }


def render_docx(context: dict) -> bytes:
    """The document, as a .docx file's bytes."""
    doc = Document()
    page(doc)

    top = doc.add_table(rows=1, cols=2)
    fixed(top, [11.8, 6])
    left, right = top.rows[0].cells
    left.width, right.width = Cm(11.8), Cm(6)
    write(left, context["store"], bold=True, size=16)
    run = left.add_paragraph().add_run("Expense report")
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED
    write(right, context["period"], bold=True, size=10.5, right=True)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    count = context["count"]
    run = p.add_run(f"{count} {'expense' if count == 1 else 'expenses'}")
    run.font.color.rgb = MUTED

    doc.add_paragraph()
    t = table(doc, ["#", "Date", "Category", "Item details", "Added by", "Amount"], WIDTHS, RIGHT)
    repeat_header(t)
    for r in context["rows"]:
        cells = row(
            t,
            [r["no"], r["date"], r["category"], r["item"], r["added_by"], r["amount"]],
            WIDTHS,
            RIGHT,
        )
        if r["note"]:
            note = cells[3].add_paragraph().add_run(r["note"])
            note.font.size = Pt(8)
            note.font.color.rgb = MUTED
    row(t, ["", "", "", "Total", "", context["total"]], WIDTHS, RIGHT, fill=TOTAL_FILL, bold=True)
    keep_rows_whole(t)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run(f"Generated {context['generated_at']} by {context['generated_by']}")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


async def expense_pdf(context: dict) -> bytes:
    """The report as a PDF, made off the event loop, sharing the production
    reports' limit on LibreOffice runs."""
    async with production_report.SLOTS:
        return await run_in_threadpool(
            lambda: production_report.docx_to_pdf(render_docx(context))
        )
